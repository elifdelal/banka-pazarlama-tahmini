"""Proje raporunu (.docx) üretir. Tablolar sonuclar/metrikler.json dosyasından okunur.

Kullanım:  python araclar/rapor_olustur.py
"""
import json
import os

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

KOK = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
M = json.load(open(os.path.join(KOK, "sonuclar", "metrikler.json"), encoding="utf-8"))

GITHUB = "https://github.com/elifdelal/banka-pazarlama-tahmini"
COLAB = "https://colab.research.google.com/drive/1z9R_0ouIuTAbumj-rWKoCLSt-kg5eJPJ?usp=sharing"
VERI_URL = "https://archive.ics.uci.edu/dataset/222/bank+marketing"

LAC = RGBColor(0x14, 0x2B, 0x4C)
GRI = RGBColor(0x52, 0x51, 0x4E)
BOYUT = 10

doc = Document()
sec = doc.sections[0]
sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
sec.top_margin, sec.bottom_margin = Cm(1.4), Cm(1.4)
sec.left_margin = sec.right_margin = Cm(1.9)
GENISLIK = 17.2  # cm, yazı alanı

stil = doc.styles["Normal"]
stil.font.name = "Calibri"
stil.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
stil.font.size = Pt(BOYUT)
stil.paragraph_format.space_after = Pt(3)
stil.paragraph_format.line_spacing = 1.05


def s(x, b=3):
    """Ondalık virgüllü sayı."""
    return f"{x:.{b}f}".replace(".", ",")


def yz(x, b=1):
    return "%" + f"{x * 100:.{b}f}".replace(".", ",")


def bin_ayir(n):
    return f"{n:,}".replace(",", ".")


def zengin(p, metin):
    """**kalın** işaretlerini destekleyen basit biçimlendirme."""
    parcalar = metin.split("**")
    for i, parca in enumerate(parcalar):
        if parca:
            r = p.add_run(parca)
            r.bold = i % 2 == 1
    return p


def baslik(metin, seviye=1):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(10 if seviye == 1 else 6)
    pf.space_after = Pt(3)
    pf.keep_with_next = True
    r = p.add_run(metin)
    r.bold = True
    r.font.size = Pt(BOYUT + (3 if seviye == 1 else 1))
    r.font.color.rgb = LAC
    if seviye == 1:  # alt çizgi
        pPr = p._p.get_or_add_pPr()
        bdr = OxmlElement("w:pBdr")
        alt = OxmlElement("w:bottom")
        for k, v in {"w:val": "single", "w:sz": "6", "w:space": "1", "w:color": "B5C4DA"}.items():
            alt.set(qn(k), v)
        bdr.append(alt)
        pPr.append(bdr)
    return p


def paragraf(metin):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    return zengin(p, metin)


def madde(metin, numarali=False):
    p = doc.add_paragraph(style="List Number" if numarali else "List Bullet")
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.left_indent = Cm(0.6)
    return zengin(p, metin)


def sekil(dosya, altyazi, genislik=GENISLIK):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(os.path.join(KOK, "grafikler", dosya), width=Cm(genislik))
    a = doc.add_paragraph()
    a.alignment = WD_ALIGN_PARAGRAPH.CENTER
    a.paragraph_format.space_after = Pt(6)
    zengin(a, altyazi)
    for r in a.runs:
        r.font.size, r.italic, r.font.color.rgb = Pt(9), True, GRI


def hucre_golge(hucre, renk):
    tcPr = hucre._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), renk)
    tcPr.append(shd)


def tablo(basliklar, satirlar, genislikler, altyazi=None, vurgu=(), sayisal=False):
    if altyazi:
        a = doc.add_paragraph()
        a.paragraph_format.space_before = Pt(4)
        a.paragraph_format.space_after = Pt(2)
        a.paragraph_format.keep_with_next = True
        zengin(a, altyazi)
        for r in a.runs:
            r.font.size, r.font.color.rgb = Pt(9), GRI
            r.italic = True
    t = doc.add_table(rows=1, cols=len(basliklar))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, b in enumerate(basliklar):
        h = t.rows[0].cells[i]
        h.text = ""
        r = h.paragraphs[0].add_run(b)
        r.bold, r.font.size, r.font.color.rgb = True, Pt(9), RGBColor(0xFF, 0xFF, 0xFF)
        hucre_golge(h, "142B4C")
    for k, satir in enumerate(satirlar):
        hucreler = t.add_row().cells
        for i, deger in enumerate(satir):
            hucreler[i].text = ""
            zengin(hucreler[i].paragraphs[0], str(deger))
            for r in hucreler[i].paragraphs[0].runs:
                r.font.size = Pt(9)
                if k in vurgu:
                    r.bold = True
            if k in vurgu:
                hucre_golge(hucreler[i], "E3ECF7")
            elif k % 2 == 1:
                hucre_golge(hucreler[i], "F5F7FA")
    # Sabit düzen: Word sütun genişliklerini içeriğe göre değiştirmesin
    t.autofit = False
    tblPr = t._tbl.tblPr
    duzen = OxmlElement("w:tblLayout")
    duzen.set(qn("w:type"), "fixed")
    tblPr.append(duzen)
    for i, g in enumerate(genislikler):
        t._tbl.tblGrid.gridCol_lst[i].w = Cm(g)
    for n, satir in enumerate(t.rows):
        trPr = satir._tr.get_or_add_trPr()
        bolme = OxmlElement("w:cantSplit")
        trPr.append(bolme)
        for i, g in enumerate(genislikler):
            satir.cells[i].width = Cm(g)
            for p in satir.cells[i].paragraphs:
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.0
                p.paragraph_format.keep_with_next = n < len(t.rows) - 1  # tablo sayfalar arasında bölünmesin
                if i > 0 and sayisal:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    bosluk = doc.add_paragraph()
    bosluk.paragraph_format.space_after = Pt(0)
    bosluk.paragraph_format.line_spacing = 0.8
    return t


def baglanti(p, url, metin=None):
    """Tıklanabilir köprü ekler."""
    rid = p.part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
                           is_external=True)
    h = OxmlElement("w:hyperlink")
    h.set(qn("r:id"), rid)
    r = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    renk = OxmlElement("w:color"); renk.set(qn("w:val"), "1F5FAD"); rpr.append(renk)
    alt = OxmlElement("w:u"); alt.set(qn("w:val"), "single"); rpr.append(alt)
    r.append(rpr)
    t = OxmlElement("w:t"); t.text = metin or url; t.set(qn("xml:space"), "preserve")
    r.append(t)
    h.append(r)
    p._p.append(h)


def sayfa_no_alt_bilgi():
    p = sec.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Elif DELAL · Banka Pazarlamasında Abonelik Tahmini · Sayfa ")
    r.font.size, r.font.color.rgb = Pt(8.5), GRI
    run = p.add_run()
    run.font.size, run.font.color.rgb = Pt(8.5), GRI
    for tip, metin in (("begin", None), (None, "PAGE"), ("end", None)):
        if tip:
            e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), tip)
        else:
            e = OxmlElement("w:instrText"); e.set(qn("xml:space"), "preserve"); e.text = metin
        run._r.append(e)


sayfa_no_alt_bilgi()

# ---------------------------------------------------------------------------
# Başlık bloğu
T, CV, V = M["test"], M["cv"], M["veri"]
hgb, lr, rf = T["HistGradientBoosting"], T["Lojistik Regresyon"], T["Random Forest"]
Z = M["zaman"]["tablo"]

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(2)
r = p.add_run("Banka Telefon Pazarlamasında Vadeli Mevduat Aboneliği Tahmini")
r.bold, r.font.size, r.font.color.rgb = True, Pt(17), LAC
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(2)
r = p.add_run("Makine Öğrenmesi ile Probleme Uygun Çözüm Geliştirme — Proje Raporu")
r.font.size, r.font.color.rgb = Pt(11.5), GRI
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(2)
r = p.add_run("Hazırlayan: Elif DELAL  ·  TNC Group Yapay Zeka Stajı  ·  Ekim 2026")
r.font.size = Pt(10)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(1)
r = p.add_run("Kod (GitHub): "); r.bold = True; r.font.size = Pt(9.5)
baglanti(p, GITHUB)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(6)
r = p.add_run("Colab not defteri: "); r.bold = True; r.font.size = Pt(9.5)
baglanti(p, COLAB)

# Özet kutusu
ozet = doc.add_table(rows=1, cols=1)
ozet.alignment = WD_TABLE_ALIGNMENT.CENTER
h = ozet.rows[0].cells[0]
h.width = Cm(GENISLIK)
hucre_golge(h, "EEF3FA")
h.text = ""
zengin(h.paragraphs[0],
       f"**Özet.** UCI Bank Marketing veri setindeki {bin_ayir(V['temiz_satir'])} telefon görüşmesi kullanılarak, "
       f"müşterinin vadeli mevduata abone olup olmayacağı **arama yapılmadan önce** tahmin edilmiştir. Pozitif sınıf "
       f"yalnızca {yz(V['pozitif_oran'])} olduğundan değerlendirme accuracy yerine precision, recall, F1 ve PR-AUC ile "
       f"yapılmıştır. Lojistik Regresyon, Random Forest ve HistGradientBoosting (HGB) karşılaştırılmış; en iyi model "
       f"HGB, test setinde **F1 = {s(hgb['F1'])}, ROC-AUC = {s(hgb['ROC-AUC'])}, PR-AUC = {s(hgb['PR-AUC'])}** "
       f"(rastgele: {s(T['Dummy']['PR-AUC'])}) elde etmiştir. Model olasılığa göre sıralanan listenin ilk %20'si "
       f"arandığında abonelerin **{yz(M['kazanc']['0.2'], 0)}**'si yakalanmaktadır. Zaman bazlı ek doğrulama, "
       f"modelin ekonomik dönem değişimlerine duyarlı olduğunu göstermiştir.")
h.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
for r in h.paragraphs[0].runs:
    r.font.size = Pt(9.5)

# ---------------------------------------------------------------------------
baslik("1. Problem Tanımı")
paragraf(
    "Bankalar vadeli mevduat gibi ürünleri telefonla pazarlarken her müşteriyi aramak ciddi zaman, personel ve "
    "müşteri memnuniyeti maliyeti doğurur; üstelik aranan müşterilerin yalnızca ~%11’i ürünü satın alır. Bu projenin "
    "**girdisi**, müşterinin demografik ve finansal bilgileri (yaş, meslek, eğitim, kredi durumu), önceki kampanyalarla "
    "ilgili geçmişi ve aramanın yapıldığı dönemin makroekonomik göstergeleridir. **Hedef çıktı**, müşterinin vadeli "
    "mevduata abone olup olmayacağıdır (ikili sınıflandırma: evet/hayır) ve bunun olasılığıdır. Gerçek dünyadaki "
    "ihtiyaç, arama listesini abone olma olasılığı en yüksek müşterilerden başlayarak sıralamak ve böylece aynı "
    "çağrı merkezi kapasitesiyle daha fazla satış yapmaktır. Model, görüşme gerçekleşmeden kullanılacağı için yalnızca "
    "arama öncesinde bilinebilen bilgilerle eğitilmiştir.")

baslik("2. Veri Seti")
paragraf(
    f"Çalışmada UCI Machine Learning Repository’de yayımlanan **Bank Marketing** veri setinin zenginleştirilmiş "
    f"sürümü (bank-additional-full.csv) kullanılmıştır (Moro, Cortez ve Rita, 2014; lisans: CC BY 4.0). Veri, bir "
    f"Portekiz bankasının Mayıs 2008 – Kasım 2010 arasında yaptığı telefon kampanyalarını tarih sırasıyla içerir: "
    f"**{bin_ayir(V['ham_satir'])} satır, 20 girdi değişkeni ve 1 hedef değişken (y)**. Kaynak: ")
baglanti(doc.paragraphs[-1], VERI_URL)
tablo(["Grup", "Değişkenler", "Tür"],
      [["Müşteri bilgisi", "age, job, marital, education, default (temerrüt), housing (konut kredisi), loan (bireysel kredi)", "1 sayısal, 6 kategorik"],
       ["Son temas", "contact (cep/sabit hat), month, day_of_week, **duration** (görüşme süresi, sn)", "1 sayısal, 3 kategorik"],
       ["Kampanya geçmişi", "campaign (bu kampanyadaki arama sayısı), pdays (son aramadan beri geçen gün), previous, poutcome", "3 sayısal, 1 kategorik"],
       ["Ekonomik bağlam", "emp.var.rate (istihdam değişimi), cons.price.idx (TÜFE), cons.conf.idx (tüketici güveni), euribor3m, nr.employed", "5 sayısal"],
       ["Hedef", "y: müşteri vadeli mevduata abone oldu mu? (yes / no)", "İkili"]],
      [3.0, 10.6, 3.4], "Tablo 1. Veri setindeki değişken grupları")
b = V["bilinmeyen"]
paragraf(
    f"**Sınıf dengesizliği:** Müşterilerin {bin_ayir(round(V['ham_satir'] * 0.11265))} tanesi (%11,3) abone olmuş, "
    f"%88,7’si olmamıştır. **Eksik değerler:** Veride boş hücre yoktur; eksik bilgiler “unknown” etiketiyle "
    f"kodlanmıştır. En çok default ({bin_ayir(b['default'])} satır, %20,9), education (%4,2), housing ve loan "
    f"(%2,4) sütunlarında görülür. Ayrıca {V['tekrar']} tekrar eden satır bulunmaktadır.")

# ---------------------------------------------------------------------------
baslik("3. Veri Analizi ve Ön İşleme")
baslik("3.1 Veri keşfi", 2)
paragraf("Kategorik ve gruplanmış sayısal değişkenler için abone olma oranları Şekil 1’de, sayısal değişkenler "
         "arasındaki ilişkiler ve görüşme süresinin etkisi Şekil 2’de verilmiştir.")
paragraf("**Gözlemler:**")
madde("**Accuracy yanıltıcıdır:** Herkese “hayır” diyen bir model %88,7 doğruluk elde eder. Bu nedenle başarı, "
      "azınlık sınıfı odaklı metriklerle (precision, recall, F1, PR-AUC) ölçülmelidir.")
madde("**Dönem etkisi çok güçlüdür:** Aramaların %33’ü mayıs ayında yapılmış ancak oran yalnızca %6’dır; mart, "
      "eylül, ekim ve aralıkta oran %44–51’e çıkar. Euribor faizi 1’in altındayken oran %46, 4’ün üstündeyken %5’tir. "
      "Bu değişkenler müşteriden çok kampanyanın yapıldığı ekonomik ortamı temsil eder (2008 krizi sonrası faizlerin düşmesi).")
madde("**Geçmiş başarı güçlü bir sinyaldir:** Önceki kampanyada ürünü almış müşterilerin %65’i yeniden abone olmuştur. "
      "Yaş ile ilişki ise U biçimlidir: 25 yaş altı (%21) ve 65 yaş üstü (%47) gruplar ortalamanın çok üzerindedir; "
      "meslekte öğrenciler ve emekliler öne çıkar. Bu doğrusal olmayan yapı ağaç tabanlı modelleri öne çıkarır.")
madde("**Çoklu doğrusallık:** emp.var.rate, euribor3m ve nr.employed arasındaki korelasyon 0,91–0,97’dir; aynı bilgiyi "
      "taşıdıkları için doğrusal modelde katsayılar tek tek yorumlanmamalıdır.")
madde("**Veri sızıntısı (data leakage):** Hedefle en yüksek korelasyonu görüşme süresi (duration, r = 0,41) taşır; abone "
      "olanların medyan görüşmesi 7,5 dk, olmayanlarınki 2,7 dk’dır. Ancak süre, arama bitmeden bilinemez ve "
      "arama bittiğinde sonuç zaten bellidir. Gerçekçi bir model için bu değişken **girdilerden çıkarılmıştır**.")

sekil("1_veri_kesfi.png", "Şekil 1. Hedef dağılımı ve değişkenlere göre abone olma oranı (kesikli çizgi: genel ortalama %11,3; n: örnek sayısı)", genislik=16.0)
sekil("2_korelasyon_sure.png", "Şekil 2. (a) Sayısal değişkenlerin korelasyon matrisi, (b) sınıflara göre görüşme süresi", genislik=14.0)
baslik("3.2 Uygulanan ön işleme adımları", 2)
tablo(["Adım", "Uygulama", "Gerekçe"],
      [["Tekrarlar", f"{V['tekrar']} tekrar eden satır silindi → {bin_ayir(V['temiz_satir'])} satır", "Aynı kaydın hem eğitim hem testte bulunmasını engellemek"],
       ["Eksik değer", "“unknown” ayrı bir kategori olarak bırakıldı", "Silmek verinin ~%25’ini kaybettirirdi; “bilinmiyor” bilgisinin kendisi anlamlı olabilir"],
       ["Sızıntı", "duration modelden çıkarıldı (yalnızca karşılaştırma için kullanıldı)", "Arama öncesinde bilinemez"],
       ["Özellik dönüşümü", "pdays = 999 (“hiç aranmadı” kodu) → ikili **onceden_arandi**", "999 gerçek bir gün sayısı değildir, ölçeklemeyi bozar"],
       ["Aykırı değer", f"campaign (maks. 56) eğitim setinin 99. yüzdeliğinde ({V['campaign_sinir']:.0f}) kırpıldı", "Aşırı sağa çarpık dağılımın doğrusal modeli bozmasını önlemek"],
       ["Encoding", f"10 kategorik değişkene One-Hot Encoding → toplam {V['ozellik_sayisi_onehot']} özellik", "Algoritmalar sayısal girdi ister; kategoriler arasında sıra yoktur"],
       ["Ölçekleme", "Sayısal değişkenlere StandardScaler (ortalama 0, std 1)", "Lojistik Regresyon’un düzenlileştirmesi ölçeğe duyarlıdır"],
       ["Veri bölme", f"%80 eğitim ({bin_ayir(V['egitim'])}) / %20 test ({bin_ayir(V['test'])}), tabakalı (stratified)", "Her iki kümede de pozitif oran %11,27 korunur"]],
      [2.6, 7.4, 7.0], "Tablo 2. Ön işleme adımları")
paragraf("Tüm dönüşümler scikit-learn **Pipeline / ColumnTransformer** içinde tanımlanmıştır; böylece ölçekleme ve "
         "encoding parametreleri yalnızca eğitim verisinden öğrenilir ve çapraz doğrulama katmanlarında bile test "
         "bilgisi modele sızmaz. Kırpma sınırı da yalnızca eğitim setinden hesaplanmıştır.")

# ---------------------------------------------------------------------------
baslik("4. Model / Yöntem Seçimi")
paragraf("Problem tablo verisi üzerinde ikili sınıflandırma olduğu için klasik makine öğrenmesi yöntemleri tercih "
         "edilmiştir; bu tür yapılandırılmış verilerde ağaç tabanlı topluluk modelleri derin öğrenmeye göre genellikle "
         "daha başarılı, daha hızlı ve daha yorumlanabilirdir. Basitten karmaşığa dört model karşılaştırılmıştır:")
tablo(["Model", "Rolü", "Seçim gerekçesi"],
      [["Dummy (çoğunluk sınıfı)", "Alt sınır", "Modellerin gerçekten öğrenip öğrenmediğini görmek için referans"],
       ["Lojistik Regresyon", "Temel (baseline) model", "Hızlı, yorumlanabilir; iyi kalibre olasılık üretir; L2 düzenlileştirme ile çoklu doğrusallığa dayanıklı"],
       ["Random Forest", "Doğrusal olmayan model", "Yüzlerce karar ağacının oylaması; U biçimli ilişkileri ve değişken etkileşimlerini yakalar, aykırı değerlere dayanıklıdır"],
       ["HistGradientBoosting", "Güçlü topluluk modeli", "Ağaçları ardışık olarak önceki hataları düzeltecek şekilde kurar; tablo verilerinde en iyi sonuçları veren yöntem ailesindendir (XGBoost/LightGBM benzeri)"]],
      [3.8, 3.4, 9.8], "Tablo 3. Karşılaştırılan modeller")
paragraf("SVM, ~33 bin satırlık eğitim verisinde çekirdek (kernel) hesaplamaları nedeniyle çok yavaş olduğundan ve doğal "
         "olasılık çıktısı vermediğinden kapsam dışında bırakılmıştır. **Dengesizlik** için sınıf ağırlıklandırma "
         "(class_weight) hiperparametre olarak denenmiş ve karar eşiği ayrıca optimize edilmiştir (Bölüm 5).")

baslik("5. Eğitim Süreci")
paragraf("Hiperparametreler yalnızca eğitim seti üzerinde **5 katlı tabakalı çapraz doğrulama** ile seçilmiştir. "
         "Lojistik Regresyon için GridSearchCV (8 kombinasyon), Random Forest (12 deneme) ve HGB (15 deneme) için "
         "RandomizedSearchCV kullanılmıştır. Seçim ölçütü **PR-AUC (average precision)**’dır; çünkü nadir pozitif sınıfı "
         "doğru sıralama başarısını doğrudan ölçer ve dengesiz veride ROC-AUC’ye göre daha ayırt edicidir. HGB’de "
         "aşırı öğrenmeyi önlemek için eğitim verisinin %10’u ile **erken durdurma** (early stopping) uygulanmıştır. "
         "Arama uzayları: LR’de C ∈ {0,01; 0,1; 1; 10}; RF’de 300 ağaç, max_depth ∈ {8, 12, 16, sınırsız}, "
         "min_samples_leaf ∈ {5, 10, 20, 40}, max_features ∈ {sqrt, 0,3}; HGB’de learning_rate ∈ {0,03; 0,05; 0,1}, "
         "max_leaf_nodes ∈ {15, 31, 63}, min_samples_leaf ∈ {20, 50, 100}, l2_regularization ∈ {0, 1}. Üç modelde de "
         "class_weight ∈ {yok, dengeli} denenmiştir.")
satirlar = []
for ad in ["Lojistik Regresyon", "Random Forest", "HistGradientBoosting"]:
    prm = CV[ad]["parametreler"]
    secilen = ", ".join(f"{k}={'yok' if v is None else (s(v, 2) if isinstance(v, float) else v)}" for k, v in prm.items())
    satirlar.append([ad, secilen, f"{s(CV[ad]['cv_pr_auc'])} ± {s(CV[ad]['cv_std'])}", s(M["esikler"][ad], 2)])
tablo(["Model", "Seçilen hiperparametreler", "CV PR-AUC (ort. ± std)", "Eşik"], satirlar,
      [3.6, 8.8, 3.0, 1.6], "Tablo 4. Hiperparametre arama sonuçları (5 katlı CV, eğitim seti)")
paragraf(
    "**Karar eşiği:** Modeller olasılık üretir ve varsayılan 0,5 eşiği dengesiz veride abonelerin çoğunu kaçırır. "
    "Eşik, test setine hiç bakılmadan, eğitim setindeki çapraz doğrulamanın katman dışı (out-of-fold) tahminleri "
    "üzerinde F1’i en yüksek yapan değer olarak seçilmiştir (≈0,17–0,21). İlginç biçimde aramada sınıf ağırlıklandırma "
    "seçilmemiştir: ağırlıklandırma olasılıkları bozarken eşik ayarı aynı etkiyi daha kontrollü sağlamaktadır.")

# ---------------------------------------------------------------------------
baslik("6. Değerlendirme Sonuçları")
sira = [("Dummy", "Dummy (çoğunluk)"), ("Lojistik Regresyon", "Lojistik Regresyon"), ("Random Forest", "Random Forest"),
        ("HistGradientBoosting", "**HistGradientBoosting**"), ("HistGradientBoosting (eşik 0,5)", "HGB – varsayılan eşik 0,5"),
        ("HGB + duration (benchmark)", "HGB + duration (gerçekçi değil)")]
satirlar = []
for anahtar, ad in sira:
    m = T[anahtar]
    satirlar.append([ad, s(m["Eşik"], 2), s(m["Accuracy"]), s(m["Precision"]), s(m["Recall"]), s(m["F1"]),
                     s(m["ROC-AUC"]), s(m["PR-AUC"])])
tablo(["Model", "Eşik", "Accuracy", "Precision", "Recall", "F1", "ROC-AUC", "PR-AUC"], satirlar,
      [4.6, 1.4, 1.8, 1.8, 1.7, 1.6, 2.0, 2.1], f"Tablo 5. Test seti sonuçları ({bin_ayir(V['test'])} müşteri, 928 abone)",
      vurgu=(3,), sayisal=True)
paragraf(
    f"Üç model de rastgele tahminin (PR-AUC {s(T['Dummy']['PR-AUC'])}) yaklaşık 4–4,5 katı PR-AUC’ye ulaşmıştır. "
    f"**HistGradientBoosting** tüm metriklerde en iyi sonucu vermiştir (F1 {s(hgb['F1'])}, PR-AUC {s(hgb['PR-AUC'])}); "
    f"Random Forest çok yakın, Lojistik Regresyon ise özellikle düşük recall bölgesinde belirgin biçimde geridedir "
    f"(Şekil 3b). Dummy modelin accuracy’si ({s(T['Dummy']['Accuracy'])}) eşik ayarlı modellerden bile yüksektir, ancak "
    f"hiçbir aboneyi bulamaz; bu, accuracy’nin bu problemde neden kullanılmaması gerektiğini açıkça gösterir. "
    f"Eşik seçiminin etkisi de büyüktür: aynı HGB modeli 0,5 eşiğinde abonelerin yalnızca {yz(T['HistGradientBoosting (eşik 0,5)']['Recall'], 0)}’ünü "
    f"yakalarken seçilen eşikte {yz(hgb['Recall'], 0)}’ini yakalar (F1 {s(T['HistGradientBoosting (eşik 0,5)']['F1'])} → {s(hgb['F1'])}).")
sekil("3_roc_pr.png", "Şekil 3. Test setinde ROC ve Precision–Recall eğrileri (noktalar: seçilen eşikteki çalışma noktası)", genislik=14.0)
cm = M["karisiklik"]
paragraf(
    f"**Karışıklık matrisi (Şekil 4a):** HGB, {bin_ayir(cm[1][0] + cm[1][1])} abonenin {cm[1][1]}’sını doğru "
    f"bulmuş, {cm[1][0]}’sini kaçırmıştır (FN). “Abone olur” dediği {bin_ayir(cm[0][1] + cm[1][1])} müşterinin {cm[0][1]}’i "
    f"(FP) abone olmamıştır; yani her ~2,1 aramadan biri satışla sonuçlanır (rastgele aramada ~8,9 aramada bir). "
    f"**Permütasyon önemine** (Şekil 4b) göre modelin en çok dayandığı değişkenler çalışan sayısı, Euribor ve "
    f"istihdam değişimi gibi ekonomik göstergeler; ardından müşterinin önceden aranmış olması, iletişim türü ve aydır. "
    f"Meslek, eğitim ve kredi bilgilerinin katkısı ise sıfıra yakındır.")
sekil("4_karisiklik_onem.png", "Şekil 4. (a) HGB karışıklık matrisi (renk: satır içi oran), (b) test setinde permütasyon önemi", genislik=14.5)
k = M["kazanc"]
paragraf(
    f"**İş açısından yorum (Şekil 5b):** Müşteriler model olasılığına göre sıralanıp listenin yalnızca ilk %10’u "
    f"arandığında abonelerin **{yz(k['0.1'], 0)}**’si, ilk %20’si arandığında **{yz(k['0.2'], 0)}**’si yakalanır "
    f"(rastgele aramaya göre {s(k['0.2'] / 0.2, 1)} kat verim). Şekil 5a ise eşiğin bir iş kararı olduğunu gösterir: "
    f"arama maliyeti yüksekse eşik yükseltilerek precision, kaçırılan müşterinin maliyeti yüksekse eşik düşürülerek "
    f"recall artırılabilir.")
sekil("5_esik_kazanc.png", "Şekil 5. (a) Eşiğe göre precision / recall / F1 değişimi, (b) kümülatif kazanç eğrisi", genislik=13.5)
bm = T["HGB + duration (benchmark)"]
paragraf(
    f"**Sızıntı karşılaştırması:** Görüşme süresi (duration) modele eklendiğinde ROC-AUC {s(hgb['ROC-AUC'])} → "
    f"{s(bm['ROC-AUC'])}, F1 {s(hgb['F1'])} → {s(bm['F1'])} sıçramaktadır. Bu yüksek değer yanıltıcıdır; bu veri setiyle sıkça "
    f"raporlanan %90+ AUC değerleri genellikle duration içeren modellere aittir. Gerçekte "
    f"kullanılabilecek model, duration içermeyen modeldir.")

# ---------------------------------------------------------------------------
baslik("7. Hata Analizi")
prof = {r["durum"]: r for r in M["hata_profil"]}
paragraf("Test setindeki tahminler doğru/yanlış türlerine ayrılarak müşteri profilleri karşılaştırılmıştır:")
tablo(["Grup", "Adet", "Ort. olasılık", "Ort. Euribor", "Önceden aranmış", "Mayıs payı", "Ort. görüşme süresi"],
      [[f"{ad} – {acik}", bin_ayir(prof[ad]["adet"]), s(prof[ad]["ort_olasilik"], 2), s(prof[ad]["euribor"], 2),
        yz(prof[ad]["onceden_arandi"], 0), yz(prof[ad]["mayis"], 0), f"{s(prof[ad]['ort_sure_dk'], 1)} dk"]
       for ad, acik in [("TP", "doğru yakalanan abone"), ("FN", "kaçırılan abone"), ("FP", "yanlış alarm"), ("TN", "doğru ret")]],
      [4.6, 1.3, 2.0, 2.0, 2.4, 2.0, 2.7], "Tablo 6. Hata türlerine göre müşteri profili (HGB, test seti)", sayisal=True)
madde(f"**Kaçırılan aboneler (FN) “kötü dönemde” abone olanlardır:** FN’lerin ortalama Euribor’u {s(prof['FN']['euribor'], 2)} "
      f"(TP’lerde {s(prof['TP']['euribor'], 2)}), %41’i mayıs ayında aranmış ve neredeyse hiçbiri önceden aranmamıştır. "
      f"Model, yüksek faiz döneminde abone olan müşterilerin yalnızca ~%4,5’ini yakalayabilmiştir (düşük faiz döneminde %82). "
      f"Bu müşterilerin ortalama görüşme süresi **{s(prof['FN']['ort_sure_dk'], 1)} dk**’dır: satışı, modelin göremediği "
      f"uzun ve ikna edici bir görüşme getirmiştir.")
madde("**Yanlış alarmlar (FP) “iyi dönemde” aranıp abone olmayanlardır:** Ortalama Euribor 1,10 ile TP’lere çok benzerler; "
      "en emin olunan FP’lerin tamamı önceki kampanyada başarılı olmuş müşterilerdir. Model, dönem ve geçmiş başarı "
      "sinyallerine aşırı güvenmekte, aynı koşullardaki müşterileri birbirinden ayıracak bireysel bilgiye (gelir, bakiye, "
      "ürün sahipliği) sahip olmadığı için ayırt edememektedir.")
zt, ze = M["zaman"]["test_oran"], M["zaman"]["egitim_oran"]
madde(f"**Zaman bazlı doğrulama:** Veri tarih sıralı olduğundan ek olarak ilk %80’lik dönemle eğitip son %20’lik dönemde "
      f"test edilmiştir. Bu dönemde abonelik oranı {yz(ze)} → {yz(zt)}’e, ortalama Euribor ~4,3 → ~1,0’a değişmiştir. "
      f"Rastgele bölmede en iyi olan HGB’nin ROC-AUC’si {s(Z['HistGradientBoosting']['ROC-AUC'])}’ye ve recall’u "
      f"{yz(Z['HistGradientBoosting']['Recall@esik'], 0)}’e düşerken, Lojistik Regresyon ROC-AUC "
      f"{s(Z['Lojistik Regresyon']['ROC-AUC'])} ve recall {yz(Z['Lojistik Regresyon']['Recall@esik'], 0)} ile en dayanıklı "
      f"model olmuştur. Ağaç modelleri eğitimde görmedikleri değer aralığına (çok düşük faiz) **dışdeğerleme "
      f"(extrapolation) yapamaz**; doğrusal model ise “faiz düştükçe abonelik artar” eğilimini sürdürebilir.")
paragraf("**Zayıf yönler:** (1) Performansın büyük kısmı müşteriden değil makroekonomik dönemden gelmektedir; "
         "(2) dönem değiştiğinde (veri kayması / concept drift) başarı ciddi biçimde düşebilir; (3) müşteri başına finansal "
         "bilgi (gelir, mevduat bakiyesi) olmadığı için aynı dönemdeki müşteriler iyi ayrıştırılamamaktadır; "
         "(4) precision ~%47 ile sınırlıdır. **Güçlü yönler:** sızıntısız ve gerçekçi kurgu, rastgele aramaya göre 3–4 kat "
         "verim, Pipeline ile tekrarlanabilir ve test seti hiç kullanılmadan seçilmiş hiperparametre ve eşik.")

# ---------------------------------------------------------------------------
baslik("8. Sonuç ve Geliştirme Önerileri")
paragraf(
    f"Bu çalışmada telefon pazarlaması öncesinde abone olacak müşterileri önceden belirleyen, sızıntıdan arındırılmış "
    f"bir sınıflandırma modeli geliştirilmiştir. HistGradientBoosting rastgele bölünmüş test setinde F1 {s(hgb['F1'])} ve "
    f"PR-AUC {s(hgb['PR-AUC'])} ile en iyi sonucu vermiş, listenin ilk %20’sini arayarak abonelerin {yz(k['0.2'], 0)}’sine "
    f"ulaşmayı mümkün kılmıştır. Ancak zaman bazlı doğrulama, ekonomik koşullar değiştiğinde daha basit Lojistik "
    f"Regresyon’un daha güvenilir olduğunu göstermiştir. Bu nedenle **uygulamada model seçimi yalnızca tek bir test "
    f"skoruna göre değil, gelecekteki dönemlerdeki dayanıklılığa göre de yapılmalıdır.** Geliştirme önerileri:")
madde("**Düzenli yeniden eğitim ve izleme:** Model her ay/çeyrek son dönem verisiyle yeniden eğitilmeli, kayan pencere "
      "(rolling window) doğrulaması kullanılmalı ve abonelik oranı ile Euribor gibi girdilerin dağılımı izlenmelidir.", True)
madde("**Dönemden bağımsız özellikler:** Ekonomik göstergeler arası yüksek korelasyon nedeniyle tek bir gösterge veya "
      "fark/değişim özellikleri (ör. son 3 aydaki faiz değişimi) kullanılmalı; müşteri bakiyesi, gelir ve ürün sahipliği "
      "gibi bireysel veriler eklenmelidir.", True)
madde("**Maliyet duyarlı eşik:** Bir aramanın maliyeti ile bir aboneliğin getirisi belirlenerek eşik F1 yerine beklenen "
      "kâra göre seçilmeli; olasılıklar kalibre edilmelidir (CalibratedClassifierCV).", True)
madde("**Daha ileri modeller ve açıklanabilirlik:** CatBoost/LightGBM ile kategorik değişkenler doğrudan işlenebilir, "
      "doğrusal ve ağaç modellerinin birleşimi (stacking) denenebilir; SHAP ile her müşteri için tahminin nedeni "
      "çağrı merkezi çalışanına gösterilebilir.", True)
madde("**Uplift modelleme:** Asıl iş sorusu “kim alır?” değil “kim aranırsa alır?” olduğu için, kontrol grubu "
      "verisi toplanarak aramanın etkisi doğrudan modellenebilir.", True)

baslik("Kaynaklar")
for kaynak in [
    "Moro, S., Cortez, P. ve Rita, P. (2014). A Data-Driven Approach to Predict the Success of Bank Telemarketing. "
    "Decision Support Systems, 62, 22–31. https://doi.org/10.1016/j.dss.2014.03.001",
    f"UCI Machine Learning Repository – Bank Marketing veri seti: {VERI_URL}",
    "Pedregosa, F. vd. (2011). Scikit-learn: Machine Learning in Python. JMLR, 12, 2825–2830.",
    f"Proje kodu: {GITHUB} · Colab not defteri: {COLAB}",
]:
    p = doc.add_paragraph(style="List Number 2")
    p.paragraph_format.space_after = Pt(1)
    r = p.add_run(kaynak)
    r.font.size = Pt(9)

cikti = os.path.join(KOK, "Proje_Raporu_Elif_DELAL.docx")
doc.save(cikti)
print("Kaydedildi:", cikti)
